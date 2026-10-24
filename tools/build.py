"""Turn [@key] citations into numbered IEEE citations (order of first appearance),
append the reference list, and build the Word file."""
import re, subprocess, sys
from pathlib import Path

sys.path.insert(0, "/home/claude/book")
from refs import REFS, PENDING

src, out_md, out_docx = sys.argv[1], sys.argv[2], sys.argv[3]
text = Path(src).read_text()
order = []
def num(m):
    keys = [k.strip().lstrip("@") for k in m.group(1).split(";")]
    nums = []
    for k in keys:
        if k not in REFS: raise SystemExit(f"unknown key {k}")
        if k not in order: order.append(k)
        nums.append(order.index(k) + 1)
    return "[" + ", ".join(map(str, nums)) + "]"
text = re.sub(r"\[(@[^\]]+)\]", num, text)
unused = set(REFS) - set(order)
if unused: print("unused refs:", unused)
text += "\n\n## References\n\n" + "\n".join(f"{i}. {REFS[k]}" for i, k in enumerate(order, 1)) + "\n"
text += "\nEntries without a DOI are proceedings that do not issue DOIs (NeurIPS, ICLR, PMLR, JMLR). Entries marked as pending cite the author's unpublished work and are completed on publication.\n"
Path(out_md).write_text(text)
print(len(order), "references")
subprocess.run(["pandoc", out_md, "-o", out_docx, "--reference-doc=/home/claude/ch01/ref.docx",
                "--resource-path=."], check=True)

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
d = Document(out_docx)
for t in d.tables:
    b = OxmlElement("w:tblBorders")
    for e in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        x = OxmlElement(f"w:{e}"); x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "4"); x.set(qn("w:color"), "999999"); b.append(x)
    t._tbl.tblPr.append(b)
    for row in t.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                for run in par.runs: run.font.size = Pt(9)
d.save(out_docx)
print("tables", len(d.tables), "images", len(d.inline_shapes))
